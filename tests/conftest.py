"""Shared pytest fixtures."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import pytest

FIXTURES = Path(__file__).parent / "fixtures"

# Point the UI's session store at a throwaway SQLite DB so importing `web.*` never
# touches the production default (/data/kb-ui). Set before any web import.
_UI_TEST_DIR = Path(tempfile.mkdtemp(prefix="kb-ui-test-"))
os.environ.setdefault("UI_DB_PATH", str(_UI_TEST_DIR / "users.db"))
os.environ.setdefault("UI_SESSION_SECRET", "test-session-secret")
# Unresolvable default host (kb-pipeline) costs a ~5 s DNS timeout per UI
# request; a closed local port fails immediately instead.
os.environ.setdefault("PIPELINE_API_URL", "http://127.0.0.1:9")
# Never load a cross-encoder in tests (stubs stand in for it).
os.environ.setdefault("RERANKER_ENABLED", "false")
# Sign-in: the fake pacod-id of tests/fake_id.py (installed per test as
# app.state.id_transport). Without one the internal URL is a closed port.
os.environ.setdefault("KB_UI_ORIGIN", "http://testserver")
os.environ.setdefault("PACOD_ID_ISSUER", "http://id.test")
os.environ.setdefault("PACOD_ID_INTERNAL_URL", "http://127.0.0.1:9")
os.environ.setdefault("PACOD_ID_UI_SECRET", "ui-client-secret-0123456789abcdef0123456789")
os.environ.setdefault("PACOD_ID_SERVICE_SECRET", "service-client-secret-0123456789abcdef0123")


@pytest.fixture(autouse=True)
def _inline_job_queue(monkeypatch: pytest.MonkeyPatch) -> None:
    """Job runners run at once in the calling thread, as BackgroundTasks did under
    TestClient (exceptions still reach the test). Production queues them on one worker."""
    from concurrent.futures import Future
    from typing import Any

    from pipeline import api

    class _Inline:
        def submit(self, fn: Any, *args: Any, **kwargs: Any) -> Future[Any]:
            future: Future[Any] = Future()
            future.set_result(fn(*args, **kwargs))
            return future

    monkeypatch.setattr(api, "_job_executor", lambda: _Inline())


@pytest.fixture()
def kb_root(tmp_path: Path) -> Path:
    """A temporary KB_ROOT with the expected sub-directories."""
    root = tmp_path / "kb"
    (root / "documents").mkdir(parents=True)
    (root / "index").mkdir()
    (root / "review").mkdir()
    return root


@pytest.fixture()
def config(kb_root: Path, monkeypatch: pytest.MonkeyPatch):  # type: ignore[type-arg]
    """Config pointing at the temporary KB_ROOT (autocommit + LLM conversion disabled)."""
    monkeypatch.setenv("KB_ROOT", str(kb_root))
    monkeypatch.setenv("QDRANT_MODE", "embedded")
    monkeypatch.setenv("GIT_AUTOCOMMIT", "false")
    monkeypatch.setenv("LLM_CONVERTER_ENABLED", "false")
    monkeypatch.delenv("QDRANT_URL", raising=False)

    from pipeline.config import load_config

    return load_config()


@pytest.fixture()
def docx_fixture(tmp_path: Path) -> Path:
    """Create a minimal DOCX fixture using python-docx."""
    pytest.importorskip("docx")
    import docx

    doc = docx.Document()
    doc.add_heading("Sample DOCX Document", level=1)
    doc.add_paragraph(
        "This is a sample DOCX document used for testing the knowledge base pipeline."
    )
    doc.add_heading("Section One", level=2)
    doc.add_paragraph(
        "Vector databases store embeddings and enable semantic similarity search. "
        "They are foundational components of modern retrieval-augmented generation systems."
    )
    doc.add_heading("Section Two", level=2)
    doc.add_paragraph(
        "Natural language processing allows computers to understand and generate human language. "
        "Applications include translation, summarisation, and question answering."
    )
    out = tmp_path / "sample.docx"
    doc.save(str(out))
    return out


@pytest.fixture()
def pdf_fixture(tmp_path: Path) -> Path:
    """Create a minimal PDF fixture using fpdf2."""
    pytest.importorskip("fpdf")
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=16)
    pdf.cell(0, 10, "Sample PDF Document", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        0,
        8,
        "This is a sample PDF document used for testing the knowledge base pipeline. "
        "It contains enough text to exercise the Docling parser.",
    )
    pdf.set_font("Helvetica", "B", size=14)
    pdf.cell(0, 10, "Section One", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(
        0,
        8,
        "Artificial intelligence encompasses machine learning, deep learning, and natural "
        "language processing. These technologies are transforming how we interact with data.",
    )
    out = tmp_path / "sample.pdf"
    pdf.output(str(out))
    return out
